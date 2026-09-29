"""Gera um SQL privado para importar uma exportação .xlsx do DG Eletros.

Uso: python3 scripts/gerar-importacao.py 'PLANILHA SITE.xlsx' importacao.sql
O SQL gerado contém dados pessoais. Nunca envie ao GitHub.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

ABAS = ('Produtos', 'Vendas', 'Usuarios', 'LogExclusoes', 'Pagamentos Comissoes', 'Horas Extras')


def sql_literal(value):
    return "'" + str(value).replace("'", "''") + "'"


def normalizar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def gerar(arquivo, destino):
    wb = load_workbook(arquivo, read_only=True, data_only=True)
    if any(nome not in wb.sheetnames for nome in ABAS):
        raise SystemExit('Faltam abas esperadas na planilha; confira se exportou a planilha correta.')
    counts = {}
    with open(destino, 'w', encoding='utf-8') as out:
        out.write('-- CONTÉM DADOS PESSOAIS. Cole no Neon. Não envie ao GitHub.\n')
        out.write('-- Execute schema.sql antes deste arquivo.\nBEGIN;\n')
        for aba in ABAS:
            rows = wb[aba].iter_rows(values_only=True)
            headers = [str(x).strip() if x is not None else '' for x in next(rows)]
            count = 0
            for row_num, row in enumerate(rows, 2):
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                record = {h: normalizar(row[i]) for i, h in enumerate(headers) if h and i < len(row) and row[i] is not None}
                if aba == 'Produtos':
                    record['ID'] = str(record.get('ID') or f'PROD-LEGACY-{row_num:05}')
                    if record.get('Quantidade') is None: record['Quantidade'] = 1
                    if not record.get('Promoção'): record['Promoção'] = 'Não'
                    key = record['ID']
                elif aba == 'Vendas': key = str(record['ID Venda'])
                elif aba == 'Usuarios':
                    key = str(record['ID'])
                    record.pop('Token', None)
                    record.pop('Validade token', None)
                elif aba == 'LogExclusoes': key = f'LOG-{row_num:05}'
                else: key = str(record['ID'])
                data = json.dumps(record, ensure_ascii=False, separators=(',', ':'))
                out.write('INSERT INTO dge_records(collection,record_key,source_row,data) VALUES ')
                out.write(f'({sql_literal(aba)},{sql_literal(key)},{row_num},{sql_literal(data)}::jsonb) ')
                out.write('ON CONFLICT (collection,record_key) DO NOTHING;\n')
                count += 1
            counts[aba] = count
        out.write('COMMIT;\n')
        for aba, count in counts.items():
            out.write(f"SELECT {sql_literal(aba)} AS aba, count(*) AS registros FROM dge_records WHERE collection={sql_literal(aba)}; -- esperado: {count}\n")
    print('SQL criado. Contagens esperadas:', ', '.join(f'{k}: {v}' for k, v in counts.items()))


if __name__ == '__main__':
    if len(sys.argv) != 3: raise SystemExit('Uso: python3 scripts/gerar-importacao.py planilha.xlsx importacao.sql')
    gerar(Path(sys.argv[1]), Path(sys.argv[2]))
